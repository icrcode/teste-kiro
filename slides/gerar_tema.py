"""
Gerador de slides SENAI sobre QUALQUER tema.

O conteúdo (textos, ícones, dados de gráficos e notas do apresentador) é escrito
pelo Claude via API; o visual é o mesmo tema SENAI de gerar_slides.py
(Open Sans, Material Symbols, azul #164194 / laranja #E84910, gráficos em HD).

Pré-requisitos:
    pip install anthropic python-pptx matplotlib pillow
    Credencial da API: variável ANTHROPIC_API_KEY  (ou `ant auth login`)

Execute:
    python slides/gerar_tema.py                                   (pergunta no terminal)
    python slides/gerar_tema.py --tema "Energia solar" --slides 10 --modo resumido
    python slides/gerar_tema.py --de-json slides/apresentacoes/energia-solar.json

Saída:
    slides/apresentacoes/<tema>.pptx  +  <tema>.json (conteúdo editável; re-renderize
    com --de-json sem nova chamada à API)
"""

import argparse
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
import unicodedata
from typing import Literal, Optional

import numpy as np
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from pydantic import BaseModel, ValidationError
from pptx import Presentation

import gerar_slides as gs
from gerar_slides import (I, X0, X1, Y0, Y1, W, H, AZUL, LARANJA, BRANCO, CINZA,
                          AZUL_BG, LAR_BG, AZUL_TXT, TINTA, SUPERF,
                          H_AZUL, H_LARANJA, H_BRANCO, H_TINTA, H_CINZA,
                          H_SERIE_1, H_SERIE_2, H_SERIE_3)

MODELO_PADRAO = "claude-opus-5"
SAIDA_DIR = os.path.join(gs.BASE, "apresentacoes")
MIN_SLIDES, MAX_SLIDES = 3, 30

LAYOUTS = ("topicos", "cartoes", "numeros", "processo", "linha_do_tempo",
           "comparacao", "grafico", "destaque", "tabela")

# Ícones Google (Material Symbols) oferecidos ao modelo — filtrados pelos codepoints
ICONES = """
lightbulb psychology school menu_book auto_stories science biotech calculate functions
analytics insights monitoring query_stats trending_up trending_down bar_chart pie_chart
show_chart timeline speed target flag rocket_launch star verified workspace_premium
emoji_objects tips_and_updates help info warning error check_circle task_alt cancel
thumb_up thumb_down favorite group groups person person_search diversity_3 handshake
support_agent record_voice_over forum chat campaign campaign mail call language public
travel_explore map location_on home apartment factory warehouse store storefront
shopping_cart payments savings account_balance attach_money currency_exchange sell
receipt_long inventory_2 local_shipping conveyor_belt precision_manufacturing construction
engineering build handyman hardware settings tune memory developer_board computer
laptop smartphone devices router dns cloud cloud_upload cloud_download database storage
dataset hub lan wifi security shield lock key vpn_key fingerprint admin_panel_settings
code terminal data_object api integration_instructions bug_report smart_toy neurology
model_training robot_2 electric_bolt bolt energy_savings_leaf solar_power wind_power
battery_charging_full electrical_services power eco nature park forest water_drop
recycling compost thermostat air co2 agriculture grass pets restaurant medical_services
health_and_safety vaccines medication monitor_heart fitness_center sports_soccer
directions_car directions_bus flight train two_wheeler electric_car traffic
schedule calendar_month event event_repeat hourglass_top history update pending
autorenew sync repeat loop swap_horiz compare_arrows alt_route route fork_right
account_tree schema category layers view_module dashboard grid_view checklist list_alt
fact_check rule gavel policy balance article description draft edit_note history_edu
quiz workspaces extension palette brush design_services photo_camera movie music_note
mic headphones videocam live_tv newspaper
""".split()

# ══════════════════════════════════════════════════════
# TEXTO QUE SE AJUSTA À CAIXA
# ══════════════════════════════════════════════════════
def _arquivo_fonte(bold, italic):
    nome = ("BoldItalic" if bold and italic else "Bold" if bold
            else "Italic" if italic else "Regular")
    return os.path.join(gs.FONT_DIR, f"OpenSans-{nome}.ttf")


def txt_fit(sl, texto, x, y, w, h, max_size, bold=False, italic=False, **kw):
    """Como gs.txt, mas reduz a fonte até o texto caber (conteúdo de tamanho livre)."""
    kw.setdefault("espaco", 1.0)
    tb = gs.txt(sl, texto, x, y, w, h, size=max_size, bold=bold, italic=italic, **kw)
    try:
        tb.text_frame.fit_text(font_family=gs.FONTE, max_size=max_size, bold=bold,
                               italic=italic, font_file=_arquivo_fonte(bold, italic))
    except Exception:
        pass   # texto vazio ou medição impossível: mantém o tamanho máximo
    return tb


def cabecalho(ctx, sl, titulo, subtitulo=""):
    gs.barra_lateral(sl)
    sl.shapes.add_picture(gs.LOGO_SENAI, I(X1 - 1.6), I(0.62), I(1.6))
    txt_fit(sl, titulo, X0, 0.55, 9.6, 0.75, 32, bold=True, anchor="m")
    if subtitulo:
        txt_fit(sl, subtitulo, X0, 1.3, 9.6, 0.42, 15, cor=CINZA)
    gs.pagina(ctx, sl)


def faixa_destaque(sl, texto, y, h=0.7, icone="lightbulb"):
    gs.rect(sl, X0, y, X1 - X0, h, fill=AZUL, raio=0.18)
    gs.icone(sl, icone, X0 + 0.3, y + (h - 0.38) / 2, 0.38, H_BRANCO)
    txt_fit(sl, texto, X0 + 0.85, y + 0.08, X1 - X0 - 1.1, h - 0.16, 15, bold=True,
            cor=BRANCO, anchor="m")


def _cores_badge(i):
    laranja = i % 2 == 1
    return (LAR_BG, H_LARANJA) if laranja else (AZUL_BG, H_AZUL)


# ══════════════════════════════════════════════════════
# GRÁFICOS GENÉRICOS (mesma linguagem visual do tema)
# ══════════════════════════════════════════════════════
def _num(v):
    """Formato numérico pt-BR: 1.234 · 12,5 · 7."""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if abs(v) >= 1000:
        return f"{v:,.0f}".replace(",", ".")
    if float(v).is_integer():
        return str(int(v))
    return f"{v:.1f}".replace(".", ",")


def _cores_series(n):
    # ordem validada para daltonismo (ver gerar_slides.py)
    return {1: [H_SERIE_2], 2: [H_SERIE_2, H_SERIE_3]}.get(n, [H_SERIE_1, H_SERIE_2, H_SERIE_3])


def _matriz(g):
    cats = g["categorias"][:12]
    series = []
    for s in g["series"][:3]:
        vals = list(s["valores"][:len(cats)])
        vals += [np.nan] * (len(cats) - len(vals))
        series.append((s["nome"], np.array(vals, float)))
    return cats, series


def g_barras(w, h, g, horizontal=False):
    cats, series = _matriz(g)
    cores = _cores_series(len(series))
    fig, ax = gs._fig(w, h)
    n = len(series); bw = 0.8 / n
    pos = np.arange(len(cats))
    rotular = len(cats) * n <= 16
    for i, ((nome, vals), c) in enumerate(zip(series, cores)):
        off = pos + (i - (n - 1) / 2) * bw
        if horizontal:
            bs = ax.barh(off, vals, bw * 0.95, label=nome, color=c, zorder=3)
        else:
            bs = ax.bar(off, vals, bw, label=nome, color=c, edgecolor=H_BRANCO,
                        linewidth=1.5, zorder=3)
        if rotular:
            ax.bar_label(bs, labels=[_num(v) for v in vals], padding=3, fontsize=9.5,
                         color=H_TINTA, fontweight="semibold")
    fmt = FuncFormatter(lambda v, _: _num(v))
    if horizontal:
        ax.set_yticks(pos, cats, color=H_TINTA); ax.invert_yaxis()
        ax.xaxis.set_major_formatter(fmt); ax.xaxis.grid(True)
        ax.spines["left"].set_visible(True)
        if g.get("eixo_y"): ax.set_xlabel(g["eixo_y"])
        ax.margins(x=0.12)
    else:
        ax.set_xticks(pos, cats)
        if max(len(c) for c in cats) * len(cats) > 70:
            plt.setp(ax.get_xticklabels(), rotation=25, ha="right")
        ax.yaxis.set_major_formatter(fmt); ax.yaxis.grid(True)
        if g.get("eixo_y"): ax.set_ylabel(g["eixo_y"])
        ax.margins(y=0.14)
    ax.set_axisbelow(True)
    if n > 1:
        ax.legend(loc="upper right", ncols=n, handlelength=1.1)
    return gs._png(fig)


def g_linhas(w, h, g):
    cats, series = _matriz(g)
    cores = _cores_series(len(series))
    x = np.arange(len(cats), dtype=float)
    fig, ax = gs._fig(w, h)
    for (nome, vals), c in zip(series, cores):
        ax.plot(*gs._suave(x, vals), color=c, lw=2.4, label=nome, zorder=3)
        ax.scatter(x, vals, s=40, color=c, edgecolor=H_BRANCO, lw=1.5, zorder=4)
        ult = np.where(~np.isnan(vals))[0]
        if len(ult) and len(series) == 1:   # com várias séries, a legenda identifica
            k = ult[-1]
            ax.annotate(_num(vals[k]), (x[k], vals[k]), xytext=(0, 9),
                        textcoords="offset points", ha="center", fontsize=10,
                        color=H_TINTA, fontweight="semibold")
    ax.set_xticks(x, cats)
    if max(len(c) for c in cats) * len(cats) > 70:
        plt.setp(ax.get_xticklabels(), rotation=25, ha="right")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: _num(v)))
    ax.yaxis.grid(True); ax.set_axisbelow(True); ax.margins(y=0.15)
    if g.get("eixo_y"): ax.set_ylabel(g["eixo_y"])
    if len(series) > 1:
        ax.legend(loc="upper left")
    return gs._png(fig)


def g_pizza(w, h, g):
    cats, series = _matriz(g)
    vals = np.nan_to_num(series[0][1])
    pares = sorted(zip(cats, vals), key=lambda p: -p[1])
    if len(pares) > 4:                       # agrupa a cauda em "Outros"
        pares = pares[:3] + [("Outros", sum(v for _, v in pares[3:]))]
    rotulos, valores = zip(*pares)
    total = sum(valores) or 1
    cores = [H_SERIE_2, H_SERIE_3, H_SERIE_1, "#B8C4D9"][:len(valores)]
    fig, ax = gs._fig(w, h)
    wedges, _ = ax.pie(valores, colors=cores, startangle=90, counterclock=False,
                       wedgeprops=dict(width=0.34, edgecolor=H_BRANCO, linewidth=3))
    for wd, lab, v in zip(wedges, rotulos, valores):
        ang = np.deg2rad((wd.theta1 + wd.theta2) / 2)
        ax.text(np.cos(ang) * 1.36, np.sin(ang) * 1.3, f"{v / total:.0%}\n{lab}",
                ha="center", va="center", fontsize=10.5, color=H_TINTA,
                fontweight="semibold", linespacing=1.15)
    ax.text(0, 0.06, f"{valores[0] / total:.0%}", ha="center", va="center",
            fontsize=22, fontweight="bold", color=H_AZUL)
    ax.text(0, -0.24, rotulos[0][:18].lower(), ha="center", va="center",
            fontsize=10, color=H_CINZA)
    ax.set_xlim(-1.8, 1.8); ax.set_ylim(-1.6, 1.6); ax.set_aspect("equal")
    return gs._png(fig)


def grafico_valido(g):
    return bool(g and g.get("categorias") and g.get("series")
                and any(len(s.get("valores", [])) for s in g["series"]))


# ══════════════════════════════════════════════════════
# LAYOUTS
# ══════════════════════════════════════════════════════
def l_topicos(ctx, sl, s):
    itens = s["itens"][:6]
    if not itens:
        return l_destaque(ctx, sl, s)
    n = len(itens)
    cols = 1 if n <= 3 else 2
    linhas = math.ceil(n / cols)
    gap = 0.22
    cw = (X1 - X0 - 0.3 * (cols - 1)) / cols
    area = Y1 - Y0 - (0.95 if s["destaque"] else 0)
    rh = min(1.8, (area - gap * (linhas - 1)) / linhas)
    for i, it in enumerate(itens):
        c, r = divmod(i, linhas)
        x, y = X0 + c * (cw + 0.3), Y0 + r * (rh + gap)
        gs.cartao(sl, x, y, cw, rh)
        d = min(0.72, rh - 0.36)
        fundo, cor = _cores_badge(i)
        gs.icone_badge(sl, it["icone"], x + 0.25, y + (rh - d) / 2, d, fundo, cor)
        tx = x + 0.25 + d + 0.28
        if it["texto"]:
            txt_fit(sl, it["titulo"], tx, y + 0.15, cw - (tx - x) - 0.25, 0.42, 17, bold=True)
            txt_fit(sl, it["texto"], tx, y + 0.6, cw - (tx - x) - 0.25, rh - 0.72, 13.5,
                    cor=CINZA, espaco=1.1)
        else:
            txt_fit(sl, it["titulo"], tx, y + 0.1, cw - (tx - x) - 0.25, rh - 0.2, 17,
                    bold=True, anchor="m")
    if s["destaque"]:
        faixa_destaque(sl, s["destaque"], Y1 - 0.7)


def l_cartoes(ctx, sl, s):
    itens = s["itens"][:4]
    if len(itens) < 2:
        return l_topicos(ctx, sl, s)
    n = len(itens)
    cw = (X1 - X0 - (n - 1) * 0.3) / n
    ch = Y1 - Y0 - (1.05 if s["destaque"] else 0)
    for i, it in enumerate(itens):
        x = X0 + i * (cw + 0.3)
        gs.cartao(sl, x, Y0, cw, ch)
        fundo, cor = _cores_badge(i)
        gs.icone_badge(sl, it["icone"], x + 0.3, Y0 + 0.32, 0.9, fundo, cor)
        txt_fit(sl, it["titulo"], x + 0.3, Y0 + 1.42, cw - 0.55, 0.75, 18, bold=True)
        txt_fit(sl, it["texto"], x + 0.3, Y0 + 2.25, cw - 0.55, ch - 2.45, 13.5,
                cor=CINZA, espaco=1.15)
    if s["destaque"]:
        faixa_destaque(sl, s["destaque"], Y1 - 0.7)


def l_numeros(ctx, sl, s):
    itens = s["itens"][:4]
    if not itens:
        return l_destaque(ctx, sl, s)
    n = len(itens)
    kw = (X1 - X0 - (n - 1) * 0.3) / n
    kh = min(3.2, Y1 - Y0 - (1.05 if s["destaque"] else 0))
    for i, it in enumerate(itens):
        x = X0 + i * (kw + 0.3)
        gs.cartao(sl, x, Y0, kw, kh)
        fundo, cor = (LAR_BG, H_LARANJA) if i == 0 else (AZUL_BG, H_AZUL)
        gs.icone_badge(sl, it["icone"], x + 0.28, Y0 + 0.28, 0.7, fundo, cor)
        txt_fit(sl, it["valor"] or "—", x + 0.28, Y0 + 1.12, kw - 0.5, 0.8, 38,
                bold=True, cor=LARANJA if i == 0 else AZUL)
        txt_fit(sl, it["titulo"], x + 0.28, Y0 + 1.95, kw - 0.5, 0.42, 14, bold=True)
        if it["texto"]:
            txt_fit(sl, it["texto"], x + 0.28, Y0 + 2.4, kw - 0.5, kh - 2.55, 12,
                    cor=CINZA, espaco=1.1)
    if s["destaque"]:
        faixa_destaque(sl, s["destaque"], Y1 - 0.7, icone="insights")


def l_processo(ctx, sl, s):
    itens = s["itens"][:6]
    if len(itens) < 2:
        return l_topicos(ctx, sl, s)
    n, gap = len(itens), 0.3
    nw = (X1 - X0 - (n - 1) * gap) / n
    ny, nh = Y0 + 0.3, 2.05
    for i, it in enumerate(itens):
        x = X0 + i * (nw + gap)
        fundo, cor = _cores_badge(i)
        gs.cartao(sl, x, ny, nw, nh)
        gs.icone_badge(sl, it["icone"], x + (nw - 0.8) / 2, ny + 0.3, 0.8, fundo, cor)
        txt_fit(sl, it["titulo"], x + 0.1, ny + 1.22, nw - 0.2, 0.7, 14.5, bold=True,
                align="c", anchor="m")
        gs.numero(sl, i + 1, x + nw - 0.32, ny - 0.14, 0.4, LARANJA if i % 2 else AZUL)
        txt_fit(sl, it["texto"], x, ny + nh + 0.15, nw, 1.3 if not s["destaque"] else 1.0,
                12.5, cor=CINZA, align="c", espaco=1.1)
        if i < n - 1:
            gs.icone(sl, "arrow_forward", x + nw + (gap - 0.26) / 2, ny + nh / 2 - 0.13,
                     0.26, H_CINZA)
    if s["destaque"]:
        faixa_destaque(sl, s["destaque"], Y1 - 0.7, icone="bolt")


def l_linha_do_tempo(ctx, sl, s):
    itens = s["itens"][:5]
    if len(itens) < 2:
        return l_topicos(ctx, sl, s)
    n = len(itens)
    cw = (X1 - X0 - (n - 1) * 0.3) / n
    ly = Y0 + 1.0
    gs.rect(sl, X0 + cw / 2, ly - 0.015, (X1 - X0) - cw, 0.03, fill=gs.LINHA)
    for i, it in enumerate(itens):
        x = X0 + i * (cw + 0.3)
        cx = x + cw / 2
        txt_fit(sl, (it["valor"] or f"Etapa {i + 1}").upper(), x, Y0, cw, 0.35, 12,
                bold=True, cor=LARANJA, align="c")
        gs.circulo(sl, cx - 0.45, ly - 0.45, 0.9, BRANCO)
        fundo, cor = (LAR_BG, H_LARANJA) if i == 0 else (AZUL_BG, H_AZUL)
        gs.icone_badge(sl, it["icone"], cx - 0.4, ly - 0.4, 0.8, fundo, cor)
        top = ly + 0.75
        gs.cartao(sl, x, top, cw, Y1 - top)
        txt_fit(sl, it["titulo"], x + 0.25, top + 0.2, cw - 0.5, 0.75, 17, bold=True)
        txt_fit(sl, it["texto"], x + 0.25, top + 1.0, cw - 0.5, Y1 - top - 1.2, 13,
                cor=CINZA, espaco=1.15)


def l_comparacao(ctx, sl, s):
    cols = s["colunas"][:2]
    if len(cols) < 2:
        return l_topicos(ctx, sl, s)
    cw = (X1 - X0 - 0.4) / 2
    base = Y1 - (0.95 if s["destaque"] else 0)
    for i, col in enumerate(cols):
        x = X0 + i * (cw + 0.4)
        cor_faixa = AZUL if i == 0 else LARANJA
        gs.rect(sl, x, Y0, cw, 0.7, fill=cor_faixa, raio=0.15)
        txt_fit(sl, col["titulo"], x + 0.3, Y0 + 0.08, cw - 0.6, 0.54, 18, bold=True,
                cor=BRANCO, anchor="m")
        top = Y0 + 0.85
        gs.cartao(sl, x, top, cw, base - top)
        itens = col["itens"][:6]
        passo = min(0.8, (base - top - 0.3) / max(len(itens), 1))
        for k, t in enumerate(itens):
            y = top + 0.2 + k * passo
            gs.icone(sl, "check_circle", x + 0.3, y + 0.04, 0.3,
                     H_AZUL if i == 0 else H_LARANJA)
            txt_fit(sl, t, x + 0.78, y, cw - 1.05, passo - 0.08, 14, cor=TINTA, espaco=1.1)
    if s["destaque"]:
        faixa_destaque(sl, s["destaque"], Y1 - 0.7, icone="compare_arrows")


def l_grafico(ctx, sl, s):
    g = s.get("grafico")
    if not grafico_valido(g):
        return l_topicos(ctx, sl, s)
    itens = s["itens"][:3]
    gw = 7.9 if itens else X1 - X0
    fn = {"barras": g_barras, "linhas": g_linhas, "pizza": g_pizza,
          "barras_horizontais": lambda w, h, d: g_barras(w, h, d, horizontal=True)}[g["tipo"]]
    gs.grafico(ctx, sl, fn, X0, Y0 - 0.1, gw, 4.45, g)
    if g.get("fonte"):
        txt_fit(sl, f"Fonte: {g['fonte']}", X0, 6.45, gw, 0.3, 10.5, italic=True, cor=CINZA)
    if itens:
        dx = X0 + gw + 0.25
        gs.cartao(sl, dx, Y0, X1 - dx, Y1 - Y0)
        passo = (Y1 - Y0 - 0.3) / len(itens)
        for i, it in enumerate(itens):
            y = Y0 + 0.25 + i * passo
            gs.icone(sl, it["icone"], dx + 0.25, y + 0.02, 0.36,
                     H_LARANJA if i == 0 else H_AZUL)
            txt_fit(sl, it["titulo"], dx + 0.75, y, X1 - dx - 1.0, 0.4, 14, bold=True)
            txt_fit(sl, it["texto"], dx + 0.75, y + 0.42, X1 - dx - 1.0, passo - 0.55, 12,
                    cor=CINZA, espaco=1.1)


def l_destaque(ctx, sl, s):
    texto = s["destaque"] or (s["itens"][0]["texto"] if s["itens"] else s["titulo"])
    gs.rect(sl, X0, Y0, X1 - X0, Y1 - Y0, fill=AZUL, raio=0.05)
    gs.icone(sl, "format_quote", X0 + 0.55, Y0 + 0.45, 0.9, H_LARANJA)
    txt_fit(sl, texto, X0 + 0.6, Y0 + 1.45, X1 - X0 - 1.2, 2.3, 32, bold=True,
            cor=BRANCO, espaco=1.1)
    if s["itens"] and s["destaque"]:
        it = s["itens"][0]
        gs.rect(sl, X0 + 0.6, Y1 - 1.05, 0.9, 0.06, fill=LARANJA)
        txt_fit(sl, f"{it['titulo']}  ·  {it['texto']}" if it["texto"] else it["titulo"],
                X0 + 0.6, Y1 - 0.9, X1 - X0 - 1.2, 0.5, 15, cor=AZUL_TXT)


def l_tabela(ctx, sl, s):
    t = s.get("tabela")
    if not t or not t.get("cabecalho") or not t.get("linhas"):
        return l_topicos(ctx, sl, s)
    cab = t["cabecalho"][:5]
    linhas = [(l + [""] * len(cab))[:len(cab)] for l in t["linhas"][:8]]
    pesos = [1.35] + [1] * (len(cab) - 1)
    total = sum(pesos)
    larg = [(X1 - X0) * p / total for p in pesos]
    hh = 0.58
    rh = min(0.8, (Y1 - Y0 - hh - 0.06) / len(linhas))
    x = X0
    gs.rect(sl, X0, Y0, X1 - X0, hh, fill=AZUL, raio=0.12)
    for c, (titulo, w) in enumerate(zip(cab, larg)):
        txt_fit(sl, titulo, x + 0.2, Y0 + 0.06, w - 0.3, hh - 0.12, 13, bold=True,
                cor=BRANCO, anchor="m")
        x += w
    for r, linha in enumerate(linhas):
        y = Y0 + hh + 0.06 + r * rh
        if r % 2 == 0:
            gs.rect(sl, X0, y, X1 - X0, rh - 0.04, fill=SUPERF, raio=0.1)
        x = X0
        for c, (cel, w) in enumerate(zip(linha, larg)):
            txt_fit(sl, cel, x + 0.2, y + 0.04, w - 0.3, rh - 0.12, 14, bold=c == 0,
                    cor=AZUL if c == 0 else TINTA, anchor="m")
            x += w


RENDER = {"topicos": l_topicos, "cartoes": l_cartoes, "numeros": l_numeros,
          "processo": l_processo, "linha_do_tempo": l_linha_do_tempo,
          "comparacao": l_comparacao, "grafico": l_grafico, "destaque": l_destaque,
          "tabela": l_tabela}


def slide_conteudo(ctx, s):
    sl = gs.novo_slide(ctx, s["notas"])
    cabecalho(ctx, sl, s["titulo"], s["subtitulo"])
    RENDER.get(s["layout"], l_topicos)(ctx, sl, s)
    gs.finalizar(ctx, sl)


def slide_capa(ctx, c, uc):
    sl = gs.novo_slide(ctx, c.get("notas_capa", ""))
    sl.shapes.add_picture(gs._logo_recortado(), I(0.75), I(0.6), height=I(1.7))
    sl.shapes.add_picture(gs.LOGO_SENAI, I(W - 0.75 - 2.2), I(0.75), I(2.2))
    gs.rect(sl, 0.75, 3.0, 1.1, 0.08, fill=LARANJA)
    txt_fit(sl, c["titulo"], 0.75, 3.2, 11.8, 1.85, 50, bold=True, espaco=0.95)
    linha = f"UC: {uc}" if uc else c["subtitulo"]
    txt_fit(sl, linha, 0.75, 5.12, 11.8, 0.55, 22, bold=True, italic=True)
    x = 0.75
    for tag in c["tags"][:5]:
        larg = min(3.2, 0.4 + 0.1 * len(tag))
        if x + larg > W - 0.75:
            break
        gs.pill(sl, tag, x, 5.85, larg, 0.38)
        x += larg + 0.15
    gs.rect(sl, 0, H - 0.12, W, 0.12, fill=AZUL)
    gs.rect(sl, 0, H - 0.12, 2.6, 0.12, fill=LARANJA)
    gs.finalizar(ctx, sl)


# ══════════════════════════════════════════════════════
# CONTEÚDO VIA CLAUDE
# ══════════════════════════════════════════════════════
def planejar(n):
    """Distribui n slides entre capa, agenda, divisórias, conteúdo e encerramento."""
    agenda = n >= 6
    secoes = 0 if n < 9 else min(4, max(2, (n - 3) // 4))
    return {"agenda": agenda, "secoes": secoes, "conteudo": n - 2 - int(agenda) - secoes}


def montar_modelos(icones):
    Icone = Literal[tuple(icones)]

    class Item(BaseModel):
        icone: Icone
        titulo: str
        texto: str
        valor: str

    class Coluna(BaseModel):
        titulo: str
        itens: list[str]

    class Serie(BaseModel):
        nome: str
        valores: list[float]

    class Grafico(BaseModel):
        tipo: Literal["barras", "barras_horizontais", "linhas", "pizza"]
        eixo_y: str
        categorias: list[str]
        series: list[Serie]
        fonte: str

    class Tabela(BaseModel):
        cabecalho: list[str]
        linhas: list[list[str]]

    class Slide(BaseModel):
        layout: Literal[LAYOUTS]
        secao: str
        titulo: str
        subtitulo: str
        itens: list[Item]
        colunas: list[Coluna]
        grafico: Optional[Grafico]
        tabela: Optional[Tabela]
        destaque: str
        notas: str

    class Secao(BaseModel):
        nome: str
        descricao: str

    class Apresentacao(BaseModel):
        titulo: str
        subtitulo: str
        tags: list[str]
        notas_capa: str
        secoes: list[Secao]
        slides: list[Slide]

    return Apresentacao


def _schema_estrito(modelo):
    """JSON Schema do Pydantic no formato aceito pelos structured outputs."""
    schema = modelo.model_json_schema()

    def ajustar(no):
        if isinstance(no, dict):
            no.pop("title", None)
            if no.get("type") == "object":
                no["additionalProperties"] = False
                no["required"] = list(no.get("properties", {}))
            for v in no.values():
                ajustar(v)
        elif isinstance(no, list):
            for v in no:
                ajustar(v)
    ajustar(schema)
    return schema


SISTEMA = """Você é designer instrucional do SENAI e cria o conteúdo de apresentações \
para cursos técnicos e profissionalizantes. Escreva em português do Brasil, com \
linguagem clara, precisa e didática, adequada ao público informado.

Cada slide usa um layout. Preencha apenas os campos que o layout usa e deixe os \
demais vazios ("" ou [] ou null):
- topicos: 3 a 6 itens (icone, titulo, texto). Explicações ou conceitos.
- cartoes: 2 a 4 itens lado a lado (icone, titulo, texto). Pilares, problemas, tipos.
- numeros: 2 a 4 itens com "valor" curto (ex.: "73%", "2,5 mi", "24h"), titulo e texto.
- processo: 3 a 6 etapas em sequência (icone, titulo, texto curto).
- linha_do_tempo: 3 a 5 marcos; "valor" é o rótulo do marco (ano, fase, prazo).
- comparacao: exatamente 2 colunas (titulo + 3 a 6 itens curtos cada).
- grafico: objeto "grafico" + até 3 itens com leituras do gráfico. Pizza usa 1 série.
- destaque: uma frase-chave forte em "destaque"; opcionalmente 1 item com a autoria \
ou contexto (titulo = nome/fonte, texto = cargo/ano).
- tabela: objeto "tabela" com 2 a 5 colunas e 2 a 8 linhas.
"destaque" também pode ser usado em topicos, cartoes, numeros, processo e \
comparacao como mensagem-resumo no rodapé do slide.

Regras:
- Varie os layouts: nunca repita o mesmo layout em slides seguidos e use pelo \
menos 4 layouts diferentes quando houver 5 ou mais slides.
- Ícones: escolha o que melhor representa cada item, só da lista permitida.
- Números e gráficos: use apenas dados que você conhece com segurança e informe a \
fonte em "fonte" (instituição e ano). Se os valores forem aproximados ou \
hipotéticos, escreva "Dados ilustrativos" na fonte. Nunca invente estatísticas \
apresentando-as como reais. Na dúvida, prefira outro layout.
- Títulos de slide com até 45 caracteres; subtítulos com até 80.
- "notas" são as notas do apresentador: o que falar naquele slide.
- "secao" de cada slide deve ser exatamente o "nome" de uma das seções listadas \
(ou "" se não houver seções). Os slides de uma mesma seção ficam juntos, na ordem \
das seções."""

NIVEIS = {
    "resumido": ("Nível RESUMIDO: textos de item com até 70 caracteres, frases diretas, "
                 "foco no essencial e nos visuais; notas com 1 a 2 frases."),
    "aprofundado": ("Nível APROFUNDADO: textos de item com até 170 caracteres, com "
                    "explicações, exemplos práticos e termos técnicos; notas com 3 a 5 "
                    "frases que aprofundam o conteúdo do slide."),
}


def montar_pedido(tema, uc, modo, plano, icones):
    secoes = (f"Organize os slides em exatamente {plano['secoes']} seções (campo "
              f"'secoes', com nome curto de até 30 caracteres e descrição de uma frase)."
              if plano["secoes"] else "Não use seções: 'secoes' deve ser [] e 'secao' \"\".")
    return (f"Tema: {tema}\n"
            f"Contexto / unidade curricular: {uc or 'não informado'}\n"
            f"{NIVEIS[modo]}\n"
            f"Crie exatamente {plano['conteudo']} slides de conteúdo (sem capa, agenda "
            f"nem encerramento — esses são gerados à parte). {secoes}\n"
            f"Ícones permitidos: {', '.join(icones)}")


def _rel(caminho):
    """Caminho relativo à raiz do projeto, com barras normais (para o chat do Kiro)."""
    return os.path.relpath(caminho, os.path.dirname(gs.BASE)).replace("\\", "/")


def exportar_para_kiro(meta, icones, caminho_json):
    """Prepara o pedido para o assistente do Kiro escrever o conteúdo.
    Grava <tema>.json com "conteudo": null e <tema>.prompt.md com as instruções."""
    with open(caminho_json, "w", encoding="utf-8") as f:
        json.dump({"meta": meta, "conteudo": None}, f, ensure_ascii=False, indent=2)
    schema = _schema_estrito(montar_modelos(icones))
    caminho_md = os.path.splitext(caminho_json)[0] + ".prompt.md"
    js, script = _rel(caminho_json), _rel(os.path.join(gs.BASE, "gerar_tema.py"))
    with open(caminho_md, "w", encoding="utf-8") as f:
        f.write(
            f"# Conteúdo de slides SENAI — {meta['tema']}\n\n"
            f"**Tarefa:** gere o conteúdo da apresentação seguindo as instruções abaixo e "
            f"salve no arquivo `{js}`, substituindo o valor `null` do campo `\"conteudo\"` "
            f"pelo objeto JSON descrito no esquema. Não altere o campo `\"meta\"`. "
            f"O arquivo deve continuar sendo JSON válido (UTF-8).\n\n"
            f"O script `{script}` fica aguardando este arquivo no terminal e monta o "
            f".pptx sozinho — não é preciso rodar nada. Se o script não estiver aberto, "
            f"monte manualmente com:\n\n"
            f"```\npython {script} --de-json {js}\n```\n\n"
            f"## Instruções\n\n{SISTEMA}\n\n"
            f"## Pedido\n\n{montar_pedido(meta['tema'], meta['uc'], meta['modo'], planejar(meta['slides']), icones)}\n\n"
            f"## Esquema JSON do campo \"conteudo\"\n\n"
            f"```json\n{json.dumps(schema, ensure_ascii=False, indent=2)}\n```\n")
    return caminho_md


def instrucoes_kiro(caminho_md, caminho_json):
    """Passo a passo manual (quando o Kiro não pode ser acionado pelo terminal)."""
    md, js = _rel(caminho_md), _rel(caminho_json)
    print(f"\n  Pedido para o Kiro salvo em:\n  {md}\n")
    print("  No chat do Kiro, envie:")
    print(f"\n    Siga as instruções de #{md}\n")
    print(f"  Depois monte a apresentação com:\n  python slides/gerar_tema.py --de-json {js}")


def _kiro_exe():
    exe = shutil.which("kiro")
    if not exe and sys.platform == "win32":
        padrao = os.path.join(os.environ.get("LOCALAPPDATA", ""), "Programs", "Kiro", "bin", "kiro.cmd")
        exe = padrao if os.path.exists(padrao) else None
    return exe


def _chamar_kiro(exe, prompt, anexo):
    """Envia o prompt ao agente do Kiro na janela já aberta do projeto."""
    cmd, env = [exe], None
    if exe.lower().endswith(".cmd"):
        # Chama o Kiro.exe direto (como o kiro.cmd faz), sem passar pelo cmd.exe,
        # para acentos e pontuação do prompt chegarem intactos.
        raiz = os.path.dirname(os.path.dirname(exe))
        exe_real = os.path.join(raiz, "Kiro.exe")
        cli = os.path.join(raiz, "resources", "app", "out", "cli.js")
        if os.path.exists(exe_real) and os.path.exists(cli):
            cmd = [exe_real, cli]
            env = dict(os.environ, ELECTRON_RUN_AS_NODE="1", VSCODE_DEV="")
    subprocess.Popen(cmd + ["chat", "-m", "agent", "-r", "-a", anexo, prompt],
                     cwd=os.path.dirname(gs.BASE), env=env,
                     stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


def _aguardar_json(caminho_json, desde, limite):
    """Espera o Kiro salvar o .json com o conteúdo preenchido (após o instante `desde`)."""
    inicio = time.time()
    while time.time() - inicio < limite:
        decorrido = int(time.time() - inicio)
        sys.stdout.write(f"\r  Aguardando o Kiro escrever o conteúdo...  {decorrido // 60:02d}:{decorrido % 60:02d}")
        sys.stdout.flush()
        time.sleep(2)
        try:
            if os.path.getmtime(caminho_json) <= desde:
                continue
            with open(caminho_json, encoding="utf-8") as f:
                dados = json.load(f)
        except (OSError, json.JSONDecodeError):
            continue   # arquivo ainda sendo escrito
        if dados.get("conteudo") is not None:
            print()
            return dados["conteudo"], os.path.getmtime(caminho_json)
    raise TimeoutError


def _rodar_kiro_cli(cli, prompt):
    """Roda o agente do Kiro CLI no terminal, sem interação, até ele terminar.
    Só ferramentas de leitura e escrita de arquivos são liberadas (nada de shell)."""
    print("  " + "─" * 66)
    r = subprocess.run([cli, "chat", "--no-interactive", "--trust-tools=read,write", prompt],
                       cwd=os.path.dirname(gs.BASE), timeout=1800)
    print("  " + "─" * 66)
    return r.returncode


def executar_com_kiro_cli(cli, meta, icones, caminho_json, tentativas=3):
    """Gera o conteúdo com o Kiro CLI e devolve o conteúdo validado."""
    caminho_md = exportar_para_kiro(meta, icones, caminho_json)
    md, js = _rel(caminho_md), _rel(caminho_json)
    modelos = montar_modelos(icones)
    prompt = (f"Leia o arquivo {md} e siga as instruções dele: escreva o conteúdo da "
              f"apresentação e salve no arquivo {js}. Não rode comandos.")
    print(f"\n  Kiro CLI escrevendo o conteúdo ({md})...\n")
    for tentativa in range(1, tentativas + 1):
        codigo = _rodar_kiro_cli(cli, prompt)
        if codigo != 0:
            print(f"\n  O Kiro CLI terminou com erro (código {codigo}).")
            if not os.environ.get("KIRO_API_KEY"):
                print("  O modo sem interação do Kiro CLI exige a variável KIRO_API_KEY "
                      "(planos Pro, Pro+, Pro Max e Power).")
            instrucoes_kiro(caminho_md, caminho_json)
            sys.exit(1)
        try:
            with open(caminho_json, encoding="utf-8") as f:
                conteudo = json.load(f).get("conteudo")
        except json.JSONDecodeError as e:
            conteudo, erros = None, f"    JSON inválido: {e}"
        else:
            erros = "    o campo \"conteudo\" continua null" if conteudo is None else ""
        if conteudo is not None:
            try:
                return modelos.model_validate(conteudo).model_dump()
            except ValidationError as e:
                erros = _erros(e)
        print(f"\n  O conteúdo tem problemas:\n{erros}")
        if tentativa == tentativas:
            sys.exit(f"\n  Corrija {js} e rode: python slides/gerar_tema.py --de-json {js}")
        print("  Pedindo ao Kiro CLI para corrigir...\n")
        linhas = " ; ".join(l.strip() for l in erros.splitlines())
        prompt = (f"O arquivo {js} tem problemas: {linhas}. Corrija seguindo as instruções "
                  f"e o esquema de {md} e salve o arquivo de novo. Não rode comandos.")


def executar_com_kiro(meta, icones, caminho_json, limite=1200, tentativas=3):
    """Gera o conteúdo pelo Kiro: CLI no terminal se instalado, senão o chat da IDE."""
    cli = shutil.which("kiro-cli")
    if cli:
        return executar_com_kiro_cli(cli, meta, icones, caminho_json, tentativas)

    caminho_md = exportar_para_kiro(meta, icones, caminho_json)
    exe = _kiro_exe()
    if not exe:
        print("\n  Kiro não encontrado no PATH.")
        instrucoes_kiro(caminho_md, caminho_json)
        sys.exit(0)
    print("\n  Kiro CLI (kiro-cli) não encontrado — usando o chat da IDE do Kiro.")

    md, js = _rel(caminho_md), _rel(caminho_json)
    modelos = montar_modelos(icones)
    desde = os.path.getmtime(caminho_json)
    prompt = (f"Siga as instruções do arquivo anexado {md}: escreva o conteúdo da "
              f"apresentação e salve no arquivo {js}. Não rode nenhum comando - o script "
              f"no terminal está aguardando e monta a apresentação automaticamente.")
    print(f"\n  Abrindo o chat do Kiro com o pedido ({md})...")
    print("  Se o Kiro pedir permissão para editar o arquivo, aprove no chat.")
    _chamar_kiro(exe, prompt, caminho_md)

    for tentativa in range(1, tentativas + 1):
        try:
            conteudo, desde = _aguardar_json(caminho_json, desde, limite)
        except TimeoutError:
            print("\n\n  Tempo esgotado esperando o Kiro.")
            instrucoes_kiro(caminho_md, caminho_json)
            sys.exit(1)
        try:
            return modelos.model_validate(conteudo).model_dump()
        except ValidationError as e:
            erros = _erros(e)
            print(f"  O conteúdo tem erros de validação:\n{erros}")
            if tentativa == tentativas:
                sys.exit(f"\n  Corrija {js} e rode: python slides/gerar_tema.py --de-json {js}")
            print("  Pedindo ao Kiro para corrigir...")
            linhas = " ; ".join(l.strip() for l in erros.splitlines())
            _chamar_kiro(exe, f"O arquivo {js} tem erros de validação: {linhas}. Corrija "
                              f"seguindo o esquema de {md} e salve o arquivo de novo. Não "
                              f"rode nenhum comando.", caminho_md)


def gerar_conteudo(tema, uc, modo, plano, modelo_id, icones):
    import anthropic

    Apresentacao = montar_modelos(icones)
    pedido = montar_pedido(tema, uc, modo, plano, icones)

    client = anthropic.Anthropic()
    print("\n  Claude está planejando o conteúdo (pode levar 1 a 3 minutos)...")
    inicio, recebidos = time.time(), 0
    with client.beta.messages.stream(
        model=modelo_id,
        max_tokens=48000,
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        thinking={"type": "adaptive"},
        output_config={"format": {"type": "json_schema",
                                  "schema": _schema_estrito(Apresentacao)}},
        system=SISTEMA,
        messages=[{"role": "user", "content": pedido}],
    ) as stream:
        for trecho in stream.text_stream:
            recebidos += len(trecho)
            sys.stdout.write(f"\r  Recebendo conteúdo: {recebidos:>6} caracteres  "
                             f"({time.time() - inicio:4.0f}s)")
            sys.stdout.flush()
        final = stream.get_final_message()
    print()

    if final.stop_reason == "refusal":
        raise RuntimeError("O modelo recusou gerar conteúdo para este tema.")
    if final.stop_reason == "max_tokens":
        raise RuntimeError("A resposta foi cortada por tamanho; tente menos slides ou o modo resumido.")
    texto = "".join(b.text for b in final.content if b.type == "text")
    return Apresentacao.model_validate_json(texto).model_dump()


# ══════════════════════════════════════════════════════
# MONTAGEM
# ══════════════════════════════════════════════════════
def montar_pptx(conteudo, meta, saida):
    plano = planejar(meta["slides"])
    slides = conteudo["slides"][:plano["conteudo"]]
    secoes = {s["nome"]: s["descricao"] for s in conteudo["secoes"]} if plano["secoes"] else {}
    if secoes:   # mantém os slides agrupados na ordem das seções
        ordem = list(secoes)
        slides.sort(key=lambda s: ordem.index(s["secao"]) if s["secao"] in ordem else len(ordem))

    roteiro = [("capa", None)]
    if plano["agenda"]:
        roteiro.append(("agenda", None))
    atual = None
    for s in slides:
        if secoes and s["secao"] != atual and s["secao"] in secoes:
            atual = s["secao"]
            roteiro.append(("divisoria", atual))
        roteiro.append(("conteudo", s))
    roteiro.append(("encerramento", None))

    prs = Presentation()
    prs.slide_width, prs.slide_height = I(W), I(H)
    ctx = gs.Ctx(prs, meta["modo"], total=len(roteiro))
    # agenda: títulos dos slides; em apresentações longas, só as seções
    usadas = [n for t, n in roteiro if t == "divisoria"]
    ctx.agenda = usadas if usadas and len(slides) > 12 else [s["titulo"] for s in slides]

    print(f"\n  Montando {len(roteiro)} slides · gráficos em {gs.DPI} dpi\n")
    n_secao = 0
    for i, (tipo, dado) in enumerate(roteiro, 1):
        rotulo = {"capa": "Capa", "agenda": "Agenda", "encerramento": "Encerramento",
                  "divisoria": f"Seção: {dado}", "conteudo": ""}[tipo] or dado["titulo"]
        gs._progresso(i - 1, len(roteiro), rotulo[:28])
        if tipo == "capa":
            slide_capa(ctx, conteudo, meta.get("uc", ""))
        elif tipo == "agenda":
            gs.slide_agenda(ctx)
        elif tipo == "divisoria":
            n_secao += 1
            gs.slide_divisoria(ctx, n_secao, dado, secoes[dado])
        elif tipo == "conteudo":
            slide_conteudo(ctx, dado)
        else:
            gs.slide_encerramento(ctx)
        gs._progresso(i, len(roteiro), rotulo[:28])
    prs.save(saida)
    return len(roteiro)


# ══════════════════════════════════════════════════════
# CLI
# ══════════════════════════════════════════════════════
def _slug(texto):
    t = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")[:60] or "apresentacao"


def _erros(e):
    """Resume erros de validação em uma linha cada (sem listar todos os ícones)."""
    linhas = []
    for err in e.errors()[:10]:
        local = ".".join(str(p) for p in err["loc"])
        msg = ("ícone fora da lista permitida" if err["type"] == "literal_error"
               and err["loc"][-1] == "icone" else err["msg"])
        linhas.append(f"    {local}: {msg} (recebido: {str(err.get('input'))[:40]!r})")
    return "\n".join(linhas)


def perguntar_tema():
    while True:
        t = gs._perguntar("\n  Tema da apresentação: ", "").strip()
        if len(t) >= 3:
            return t
        print("  Informe um tema com pelo menos 3 caracteres.")


def perguntar_slides():
    while True:
        r = gs._perguntar(f"\n  Quantos slides? ({MIN_SLIDES} a {MAX_SLIDES}) [10]: ", 10)
        if r.isdigit() and MIN_SLIDES <= int(r) <= MAX_SLIDES:
            return int(r)
        print(f"  Digite um número inteiro entre {MIN_SLIDES} e {MAX_SLIDES}.")


def main():
    for fluxo in (sys.stdout, sys.stderr):
        try:
            fluxo.reconfigure(encoding="utf-8", errors="replace")
        except AttributeError:
            pass

    ap = argparse.ArgumentParser(description="Gera apresentação SENAI sobre qualquer tema (conteúdo via Claude).")
    ap.add_argument("--tema", help="tema da apresentação")
    ap.add_argument("--uc", default=None, help="unidade curricular ou contexto (opcional)")
    ap.add_argument("--modo", choices=["resumido", "aprofundado"])
    ap.add_argument("--slides", type=int, help=f"quantidade de slides ({MIN_SLIDES} a {MAX_SLIDES})")
    ap.add_argument("--de-json", help="renderiza um conteúdo salvo, sem chamar a API")
    ap.add_argument("--kiro", action="store_true",
                    help="o assistente do Kiro escreve o conteúdo (não usa chave de API)")
    ap.add_argument("--modelo", default=MODELO_PADRAO, help=f"modelo Claude (padrão {MODELO_PADRAO})")
    ap.add_argument("--dpi", type=int, default=gs.DPI, help="resolução dos gráficos (padrão 300)")
    ap.add_argument("--saida", help="caminho do .pptx (padrão slides/apresentacoes/<tema>.pptx)")
    ap.add_argument("--sem-instalar-fontes", action="store_true",
                    help="não instala Open Sans no Windows do usuário")
    args = ap.parse_args()
    gs.DPI = args.dpi
    if args.slides is not None and not MIN_SLIDES <= args.slides <= MAX_SLIDES:
        ap.error(f"--slides deve estar entre {MIN_SLIDES} e {MAX_SLIDES}")

    print("\n  SENAI  |  Gerador de slides sobre qualquer tema")
    print("  " + "─" * 66)

    print("\n  Preparando recursos (Google Fonts, ícones, logos)...")
    gs.garantir_fontes()
    if not args.sem_instalar_fontes:
        novas = gs.instalar_fontes_usuario()
        if novas:
            print(f"  Open Sans instalada para o usuário ({len(novas)} arquivos).")
    with open(gs.SYMBOLS_CP, encoding="utf-8") as f:
        existentes = {l.split()[0] for l in f if l.strip()}
    icones = sorted({i for i in ICONES if i in existentes})

    if args.de_json:
        with open(args.de_json, encoding="utf-8") as f:
            dados = json.load(f)
        meta = dados["meta"]
        if dados.get("conteudo") is None:
            md = os.path.splitext(args.de_json)[0] + ".prompt.md"
            sys.exit(f"\n  O conteúdo ainda não foi preenchido. Peça ao Kiro para seguir "
                     f"as instruções de\n  {_rel(md)}")
        if args.slides:
            meta["slides"] = args.slides
        try:   # valida o JSON editado à mão contra o mesmo esquema
            conteudo = montar_modelos(icones).model_validate(dados["conteudo"]).model_dump()
        except ValidationError as e:
            sys.exit(f"\n  JSON inválido:\n{_erros(e)}")
        saida = args.saida or os.path.splitext(args.de_json)[0] + ".pptx"
    else:
        tema = args.tema or perguntar_tema()
        uc = args.uc if args.uc is not None else gs._perguntar(
            "\n  Unidade curricular ou contexto (opcional, Enter para pular): ", "").strip()
        modo = args.modo or gs.perguntar_modo()
        n = args.slides or perguntar_slides()
        meta = {"tema": tema, "uc": uc, "modo": modo, "slides": n}

        os.makedirs(SAIDA_DIR, exist_ok=True)
        saida = args.saida or os.path.join(SAIDA_DIR, _slug(tema) + ".pptx")
        caminho_json = os.path.splitext(saida)[0] + ".json"

        usar_kiro = args.kiro or not (os.environ.get("ANTHROPIC_API_KEY")
                                      or os.environ.get("ANTHROPIC_AUTH_TOKEN"))
        if usar_kiro:
            if not args.kiro:
                print("\n  Nenhuma chave da API encontrada (ANTHROPIC_API_KEY) — "
                      "o Kiro vai escrever o conteúdo.")
            try:
                conteudo = executar_com_kiro(meta, icones, caminho_json)
            except KeyboardInterrupt:
                sys.exit(f"\n\n  Interrompido. Quando o Kiro terminar, rode:\n"
                         f"  python slides/gerar_tema.py --de-json {_rel(caminho_json)}")
            total = montar_pptx(conteudo, meta, saida)
            print(f"\n\n  Concluído: {total} slides. Arquivo salvo em:\n  {saida}\n")
            return

        import anthropic
        try:
            conteudo = gerar_conteudo(tema, uc, modo, planejar(n), args.modelo, icones)
        except anthropic.AuthenticationError:
            sys.exit("\n  Credencial da API inválida. Confira ANTHROPIC_API_KEY "
                     "ou use --kiro para o Kiro escrever o conteúdo.")
        except anthropic.RateLimitError:
            sys.exit("\n  Limite de uso da API atingido. Aguarde alguns instantes e tente de novo.")
        except anthropic.APIStatusError as e:
            sys.exit(f"\n  Erro da API ({e.status_code}): {e.message}")
        except anthropic.APIConnectionError:
            sys.exit("\n  Sem conexão com a API. Verifique a internet.")
        except ValidationError as e:
            sys.exit(f"\n  A resposta do modelo não seguiu o esquema:\n{_erros(e)}")
        except RuntimeError as e:
            sys.exit(f"\n  {e}")

        meta["modelo"] = args.modelo
        with open(caminho_json, "w", encoding="utf-8") as f:
            json.dump({"meta": meta, "conteudo": conteudo}, f, ensure_ascii=False, indent=2)
        print(f"  Conteúdo salvo em {caminho_json}")

    total = montar_pptx(conteudo, meta, saida)
    print(f"\n\n  Concluído: {total} slides. Arquivo salvo em:\n  {saida}\n")


if __name__ == "__main__":
    main()
