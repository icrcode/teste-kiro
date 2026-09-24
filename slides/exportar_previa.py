"""
Exporta os slides de um .pptx gerado por gerar_slides.py / gerar_tema.py como imagens PNG.

Não depende do PowerPoint nem do LibreOffice: desenha formas, imagens e textos
(Open Sans) com o Pillow. É fiel ao que os geradores produzem (retângulos, círculos,
polígonos, imagens e caixas de texto); outros recursos do PowerPoint não são suportados.

Execute:
    python slides/exportar_previa.py resultado-slide/candelabro.pptx
    python slides/exportar_previa.py arquivo.pptx --largura 1600 --saida pasta/

Saída padrão: <pasta do pptx>/previa/<nome>/slide-01.png ... + grade.png (visão geral)
"""

import argparse
import io
import os
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

FONT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "fonts")
EMU = 914400
ESCALA = 2            # supersampling para bordas e textos suaves
NS_A = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

_fontes = {}

def _fonte(bold, italic, px):
    nome = ("BoldItalic" if bold and italic else "Bold" if bold
            else "Italic" if italic else "Regular")
    chave = (nome, int(px))
    if chave not in _fontes:
        _fontes[chave] = ImageFont.truetype(os.path.join(FONT_DIR, f"OpenSans-{nome}.ttf"),
                                            max(1, int(px)))
    return _fontes[chave]


def _cor(fill):
    try:
        if fill.type == 1:          # MSO_FILL.SOLID
            return "#" + str(fill.fore_color.rgb)
    except (AttributeError, TypeError):
        pass
    return None


def _quebrar(draw, texto, fonte, largura):
    linhas = []
    for paragrafo in texto.split("\n"):
        atual = ""
        for palavra in paragrafo.split(" "):
            teste = f"{atual} {palavra}".strip()
            if not atual or draw.textlength(teste, font=fonte) <= largura:
                atual = teste
            else:
                linhas.append(atual)
                atual = palavra
        linhas.append(atual)
    return linhas


def _texto(draw, forma, px):
    tf = forma.text_frame
    x, y, w, h = (px(forma.left), px(forma.top), px(forma.width), px(forma.height))
    ml, mr = px(tf.margin_left or 0), px(tf.margin_right or 0)
    mt, mb = px(tf.margin_top or 0), px(tf.margin_bottom or 0)
    linhas = []
    for par in tf.paragraphs:
        if not par.runs:
            continue
        run = par.runs[0]
        tam = (run.font.size.pt if run.font.size else 18) / 72 * px(EMU)
        fonte = _fonte(bool(run.font.bold), bool(run.font.italic), tam)
        cor = ("#" + str(run.font.color.rgb)
               if run.font.color and run.font.color.type is not None else "#000000")
        esp = par.line_spacing if isinstance(par.line_spacing, float) else 1.0
        texto = "".join(r.text for r in par.runs)
        for linha in _quebrar(draw, texto, fonte, w - ml - mr + 1):
            linhas.append((linha, fonte, cor, tam * 1.36 * esp, par.alignment))
    altura = sum(l[3] for l in linhas)
    ancora = tf.vertical_anchor
    if ancora == MSO_ANCHOR.MIDDLE or (ancora is None and forma.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE):
        cy = y + (h - altura) / 2
    elif ancora == MSO_ANCHOR.BOTTOM:
        cy = y + h - mb - altura
    else:
        cy = y + mt
    for linha, fonte, cor, alt, alinhamento in linhas:
        lt = draw.textlength(linha, font=fonte)
        if alinhamento == PP_ALIGN.CENTER:
            cx = x + (w - lt) / 2
        elif alinhamento == PP_ALIGN.RIGHT:
            cx = x + w - mr - lt
        else:
            cx = x + ml
        draw.text((cx, cy + alt * 0.08), linha, font=fonte, fill=cor)
        cy += alt


def renderizar(prs, slide, largura):
    ppi = largura / (prs.slide_width / EMU) * ESCALA
    px = lambda v: v / EMU * ppi
    W, H = int(px(prs.slide_width)), int(px(prs.slide_height))
    im = Image.new("RGB", (W, H), _cor(slide.background.fill) or "#FFFFFF")
    draw = ImageDraw.Draw(im)
    for forma in slide.shapes:
        x, y, w, h = px(forma.left), px(forma.top), px(forma.width), px(forma.height)
        tipo = forma.shape_type
        if tipo == MSO_SHAPE_TYPE.PICTURE:
            pic = Image.open(io.BytesIO(forma.image.blob)).convert("RGBA")
            pic = pic.resize((max(1, round(w)), max(1, round(h))), Image.LANCZOS)
            im.paste(pic, (round(x), round(y)), pic)
        elif tipo == MSO_SHAPE_TYPE.FREEFORM:
            pontos = [(x + int(p.get("x")) / EMU * ppi, y + int(p.get("y")) / EMU * ppi)
                      for p in forma._element.findall(".//a:pt", NS_A)]
            cor = _cor(forma.fill)
            if pontos and cor:
                draw.polygon(pontos, fill=cor)
        elif tipo == MSO_SHAPE_TYPE.AUTO_SHAPE:
            cor = _cor(forma.fill)
            if cor:
                caixa = [x, y, x + w, y + h]
                estilo = forma.auto_shape_type
                if estilo == MSO_SHAPE.OVAL:
                    draw.ellipse(caixa, fill=cor)
                elif estilo == MSO_SHAPE.ROUNDED_RECTANGLE:
                    draw.rounded_rectangle(caixa, radius=forma.adjustments[0] * min(w, h), fill=cor)
                else:
                    draw.rectangle(caixa, fill=cor)
            if forma.has_text_frame and forma.text_frame.text:
                _texto(draw, forma, px)
        if tipo == MSO_SHAPE_TYPE.TEXT_BOX:
            _texto(draw, forma, px)
    return im.resize((W // ESCALA, H // ESCALA), Image.LANCZOS)


def grade(imagens, colunas=3, largura=1800, margem=16):
    """Monta uma visão geral (contact sheet) dos slides."""
    cel = (largura - margem * (colunas + 1)) // colunas
    alt = round(cel * imagens[0].height / imagens[0].width)
    linhas = -(-len(imagens) // colunas)
    folha = Image.new("RGB", (largura, linhas * (alt + margem) + margem), "#E6ECF7")
    for i, im in enumerate(imagens):
        c, r = i % colunas, i // colunas
        folha.paste(im.resize((cel, alt), Image.LANCZOS),
                    (margem + c * (cel + margem), margem + r * (alt + margem)))
    return folha


def main():
    ap = argparse.ArgumentParser(description="Exporta slides .pptx como PNG.")
    ap.add_argument("pptx")
    ap.add_argument("--largura", type=int, default=1280, help="largura de cada imagem em px")
    ap.add_argument("--saida", help="pasta de saída")
    args = ap.parse_args()

    nome = os.path.splitext(os.path.basename(args.pptx))[0]
    saida = args.saida or os.path.join(os.path.dirname(os.path.abspath(args.pptx)), "previa", nome)
    os.makedirs(saida, exist_ok=True)

    prs = Presentation(args.pptx)
    imagens = []
    for i, slide in enumerate(prs.slides, 1):
        im = renderizar(prs, slide, args.largura)
        im.save(os.path.join(saida, f"slide-{i:02d}.png"), optimize=True)
        imagens.append(im)
        sys.stdout.write(f"\r  {nome}: {i}/{len(prs.slides)} slides")
        sys.stdout.flush()
    grade(imagens).save(os.path.join(saida, "grade.png"), optimize=True)
    print(f"\n  Imagens em {saida}")


if __name__ == "__main__":
    main()
