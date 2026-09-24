"""
Gerador de slides PowerPoint — Previsão de Demanda de Vagas por Curso
Execute:  python slides/gerar_slides.py
Saída:    slides/previsao-matriculas.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.oxml.ns import qn
from lxml import etree
import copy
import os

# ---------------------------------------------------------------------------
# Paleta de cores
# ---------------------------------------------------------------------------
C_BG_DARK    = RGBColor(0x0D, 0x1B, 0x2A)   # azul-marinho fundo
C_BG_MID     = RGBColor(0x17, 0x2A, 0x3F)   # azul medio painel
C_ACCENT     = RGBColor(0x00, 0xB4, 0xD8)   # cyan AWS
C_ACCENT2    = RGBColor(0xFF, 0x9F, 0x1C)   # laranja destaque
C_WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT      = RGBColor(0xCC, 0xE8, 0xF4)   # azul claro texto
C_GRAY       = RGBColor(0xA0, 0xB4, 0xC8)
C_GREEN      = RGBColor(0x06, 0xD6, 0xA0)
C_PURPLE     = RGBColor(0x9B, 0x5D, 0xE5)

W = Inches(13.33)   # widescreen 16:9
H = Inches(7.5)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def rgb_hex(r):
    return f"{r[0]:02X}{r[1]:02X}{r[2]:02X}"

def set_bg(slide, color):
    """Preenche o fundo do slide com uma cor sólida."""
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, x, y, w, h, fill=None, alpha_emu=None, line=None, line_w=Pt(0)):
    from pptx.util import Pt
    shape = slide.shapes.add_shape(1, x, y, w, h)   # MSO_SHAPE.RECTANGLE
    shape.line.fill.background()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = line_w
    else:
        shape.line.fill.background()
    return shape

def add_label(slide, text, x, y, w, h,
              font_size=Pt(14), bold=False, color=C_WHITE,
              align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txb = slide.shapes.add_textbox(x, y, w, h)
    tf  = txb.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size   = font_size
    run.font.bold   = bold
    run.font.color.rgb = color
    run.font.italic = italic
    return txb

def gradient_rect(slide, x, y, w, h, c1: RGBColor, c2: RGBColor, angle=0):
    """Retângulo com gradiente linear usando XML direto."""
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.line.fill.background()
    sp = shape._element
    spPr = sp.find(qn("p:spPr"))
    # remove fill simples
    solidFill = spPr.find(qn("a:solidFill"))
    if solidFill is not None:
        spPr.remove(solidFill)
    gradFill_xml = f"""
<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">
  <a:gsLst>
    <a:gs pos="0">
      <a:srgbClr val="{rgb_hex(c1)}"/>
    </a:gs>
    <a:gs pos="100000">
      <a:srgbClr val="{rgb_hex(c2)}"/>
    </a:gs>
  </a:gsLst>
  <a:lin ang="{angle}" scaled="0"/>
</a:gradFill>"""
    spPr.append(etree.fromstring(gradFill_xml))
    return shape

def add_pill(slide, text, x, y, w, h, bg, fg=C_WHITE, font_size=Pt(11), bold=True):
    """Pill/badge arredondado."""
    shape = slide.shapes.add_shape(5, x, y, w, h)   # ROUNDED_RECTANGLE
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    tf = shape.text_frame
    tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size  = font_size
    run.font.bold  = bold
    run.font.color.rgb = fg
    return shape

def add_divider(slide, y, color=C_ACCENT, alpha=0.4):
    add_rect(slide, Inches(0.8), y, Inches(11.73), Pt(1.5), fill=color)

# ---------------------------------------------------------------------------
# Slide 1 — Capa
# ---------------------------------------------------------------------------

def slide_capa(prs):
    layout = prs.slide_layouts[6]   # blank
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)

    # Faixa de gradiente no rodapé
    gradient_rect(sl, 0, Inches(5.8), W, Inches(1.7), C_BG_MID, C_BG_DARK, angle=5400000)

    # Linha decorativa topo
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    # Linha decorativa fundo
    add_rect(sl, 0, Inches(7.42), W, Inches(0.08), fill=C_ACCENT2)

    # Badge AWS
    add_pill(sl, "☁  AWS  ·  Machine Learning", Inches(0.8), Inches(1.2),
             Inches(3.6), Inches(0.45), C_ACCENT, font_size=Pt(12))

    # Título principal
    add_label(sl, "Previsão de Demanda", Inches(0.8), Inches(1.9),
              Inches(11.5), Inches(1.1), font_size=Pt(52), bold=True,
              color=C_WHITE, align=PP_ALIGN.LEFT)
    add_label(sl, "de Vagas por Curso", Inches(0.8), Inches(2.85),
              Inches(11.5), Inches(1.1), font_size=Pt(52), bold=True,
              color=C_ACCENT, align=PP_ALIGN.LEFT)

    # Subtítulo
    add_label(sl,
              "Pipeline AWS end-to-end · XGBoost · SageMaker · QuickSight",
              Inches(0.8), Inches(4.0), Inches(10), Inches(0.55),
              font_size=Pt(18), color=C_LIGHT, align=PP_ALIGN.LEFT)

    # Rodapé
    add_label(sl, "Workshop de Arquiteturas AWS  ·  Kiro IDE",
              Inches(0.8), Inches(6.85), Inches(8), Inches(0.4),
              font_size=Pt(11), color=C_GRAY, align=PP_ALIGN.LEFT)
    add_label(sl, "2026", Inches(11.5), Inches(6.85), Inches(1.5), Inches(0.4),
              font_size=Pt(11), color=C_GRAY, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------------------
# Slide 2 — Problema e Solução
# ---------------------------------------------------------------------------

def slide_problema(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "O Problema", Inches(0.7), Inches(0.25),
              Inches(6), Inches(0.7), font_size=Pt(28), bold=True,
              color=C_WHITE)
    add_divider(sl, Inches(0.95))

    problemas = [
        ("📉", "Superdimensionamento de vagas", "Custo elevado e salas ociosas"),
        ("📈", "Subdimensionamento de vagas",   "Fila de espera e frustração"),
        ("🔍", "Decisão baseada em intuição",   "Sem evidências históricas"),
        ("🗓", "Planejamento reativo",           "Sem antecipação semestral"),
    ]

    for i, (icon, title, sub) in enumerate(problemas):
        y = Inches(1.15) + i * Inches(1.38)
        # card
        add_rect(sl, Inches(0.65), y, Inches(5.8), Inches(1.2),
                 fill=C_BG_MID, line=C_ACCENT2, line_w=Pt(1))
        add_label(sl, icon,  Inches(0.85), y + Inches(0.28), Inches(0.6), Inches(0.6),
                  font_size=Pt(22))
        add_label(sl, title, Inches(1.5), y + Inches(0.1), Inches(4.8), Inches(0.5),
                  font_size=Pt(14), bold=True, color=C_WHITE)
        add_label(sl, sub,   Inches(1.5), y + Inches(0.58), Inches(4.8), Inches(0.45),
                  font_size=Pt(11), color=C_GRAY)

    # Solução (direita)
    add_rect(sl, Inches(7.0), Inches(1.05), Inches(5.8), Inches(5.9),
             fill=C_BG_MID, line=C_GREEN, line_w=Pt(2))
    add_label(sl, "✅  A Solução", Inches(7.2), Inches(1.15), Inches(5.4), Inches(0.55),
              font_size=Pt(18), bold=True, color=C_GREEN)

    solucao = [
        "Modelo ML treinado com histórico de matrículas",
        "Horizonte de previsão: 2 semestres à frente",
        "Dados: ENEM, perfil socioeconômico, calendário",
        "Algoritmo principal: XGBoost",
        "Retreinamento automático a cada semestre",
        "Dashboard para gestores acadêmicos",
    ]
    for i, linha in enumerate(solucao):
        add_label(sl, f"→  {linha}",
                  Inches(7.3), Inches(1.85) + i * Inches(0.78),
                  Inches(5.3), Inches(0.65),
                  font_size=Pt(13), color=C_LIGHT)

# ---------------------------------------------------------------------------
# Slide 3 — Fontes de Dados
# ---------------------------------------------------------------------------

def slide_dados(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Fontes de Dados", Inches(0.7), Inches(0.25),
              Inches(9), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    fontes = [
        ("📚", "Histórico de Matrículas",  "Sistema Acadêmico (SIA)",
         "Séries históricas · CSV/Parquet · Freq. semestral", C_ACCENT),
        ("📋", "Catálogo de Cursos",        "Pró-Reitoria de Graduação",
         "Modalidades, capacidades · CSV", C_ACCENT2),
        ("👤", "Perfil dos Alunos",         "ENEM / Vestibular",
         "Dados demográficos e socioeconômicos · JSON/CSV", C_GREEN),
        ("📅", "Calendário Acadêmico",      "Secretaria Acadêmica",
         "Datas de inscrições e períodos letivos · JSON", C_PURPLE),
    ]

    for i, (icon, nome, fonte, desc, cor) in enumerate(fontes):
        col = i % 2
        row = i // 2
        x = Inches(0.65) + col * Inches(6.2)
        y = Inches(1.2) + row * Inches(2.7)
        add_rect(sl, x, y, Inches(5.9), Inches(2.4),
                 fill=C_BG_MID, line=cor, line_w=Pt(2))
        add_label(sl, icon, x + Inches(0.15), y + Inches(0.1),
                  Inches(0.7), Inches(0.7), font_size=Pt(28))
        add_label(sl, nome, x + Inches(0.9), y + Inches(0.1),
                  Inches(4.7), Inches(0.5), font_size=Pt(14), bold=True, color=C_WHITE)
        add_label(sl, fonte, x + Inches(0.9), y + Inches(0.58),
                  Inches(4.7), Inches(0.35), font_size=Pt(11), color=cor, italic=True)
        add_label(sl, desc, x + Inches(0.15), y + Inches(1.05),
                  Inches(5.5), Inches(1.1), font_size=Pt(11), color=C_LIGHT)

# ---------------------------------------------------------------------------
# Slide 4 — Features do Modelo
# ---------------------------------------------------------------------------

def slide_features(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Feature Engineering", Inches(0.7), Inches(0.25),
              Inches(9), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))
    add_label(sl, "9 features selecionadas para o modelo XGBoost",
              Inches(0.7), Inches(1.0), Inches(11), Inches(0.4),
              font_size=Pt(14), color=C_GRAY)

    features = [
        ("🆔", "curso_id",                        "Identificador único do curso",              C_ACCENT),
        ("📆", "periodo_letivo",                   "Semestre/ano do período previsto",          C_ACCENT),
        ("📊", "historico_demanda_3anos",          "Média de matrículas dos últimos 3 anos",    C_ACCENT2),
        ("📉", "taxa_evasao_historica",            "% histórico de evasão por curso",           C_ACCENT2),
        ("🎯", "vagas_ofertadas_periodo_anterior", "Vagas do semestre imediatamente anterior",  C_GREEN),
        ("📝", "nota_corte_enem",                  "Nota de corte no ENEM por curso",           C_GREEN),
        ("🗺", "regiao_geografica",                "Macroregião da instituição",                C_PURPLE),
        ("💻", "modalidade_ensino",                "Presencial · EAD · Híbrido",               C_PURPLE),
        ("🏙", "IDH_municipio",                   "Índice de desenvolvimento humano local",    C_ACCENT2),
    ]

    cols = 3
    for i, (icon, feat, desc, cor) in enumerate(features):
        col = i % cols
        row = i // cols
        x = Inches(0.55) + col * Inches(4.22)
        y = Inches(1.55) + row * Inches(1.8)
        add_rect(sl, x, y, Inches(4.0), Inches(1.6),
                 fill=C_BG_MID, line=cor, line_w=Pt(1.5))
        add_label(sl, icon, x + Inches(0.12), y + Inches(0.12),
                  Inches(0.5), Inches(0.5), font_size=Pt(18))
        add_label(sl, feat, x + Inches(0.65), y + Inches(0.08),
                  Inches(3.2), Inches(0.5), font_size=Pt(11), bold=True, color=cor)
        add_label(sl, desc, x + Inches(0.12), y + Inches(0.72),
                  Inches(3.7), Inches(0.7), font_size=Pt(10), color=C_LIGHT)

# ---------------------------------------------------------------------------
# Slide 5 — Arquitetura AWS
# ---------------------------------------------------------------------------

def slide_arquitetura(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Arquitetura AWS", Inches(0.7), Inches(0.25),
              Inches(9), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    # Camadas
    camadas = [
        ("INGESTÃO",      C_ACCENT2,  Inches(0.5),  Inches(1.1)),
        ("ETL / FEATURES",C_ACCENT,   Inches(0.5),  Inches(2.4)),
        ("ML PIPELINE",   C_GREEN,    Inches(0.5),  Inches(3.7)),
        ("OPS / MONITOR", C_PURPLE,   Inches(0.5),  Inches(5.0)),
        ("VISUALIZAÇÃO",  C_ACCENT,   Inches(0.5),  Inches(6.3)),
    ]
    for label, cor, x, y in camadas:
        add_rect(sl, x, y, Inches(1.7), Inches(1.05),
                 fill=None, line=cor, line_w=Pt(1))
        add_label(sl, label, x + Inches(0.05), y + Inches(0.28),
                  Inches(1.6), Inches(0.5), font_size=Pt(9), bold=True,
                  color=cor, align=PP_ALIGN.CENTER)

    # Serviços por camada
    servicos = [
        # (camada_y_base, lista de (nome, tipo, cor_badge))
        (Inches(1.1), [
            ("S3 Raw\n(prevmatriculas-dev-raw)", "S3", C_ACCENT2),
        ]),
        (Inches(2.4), [
            ("Glue ETL\nFeature Engineering", "Glue", C_ACCENT),
            ("S3 Processed\n(prevmatriculas-dev-processed)", "S3", C_ACCENT),
        ]),
        (Inches(3.7), [
            ("SageMaker\nFeature Store", "SM", C_GREEN),
            ("SageMaker\nTraining (XGBoost)", "SM", C_GREEN),
            ("SageMaker\nEndpoint (Inferência)", "SM", C_GREEN),
            ("Lambda\nPré-processador", "λ", C_GREEN),
        ]),
        (Inches(5.0), [
            ("SageMaker\nModel Monitor", "SM", C_PURPLE),
            ("CloudWatch\nMétricas", "CW", C_PURPLE),
            ("EventBridge\nAgendador", "EB", C_PURPLE),
        ]),
        (Inches(6.3), [
            ("QuickSight\nDashboard Gestores", "QS", C_ACCENT),
        ]),
    ]

    for (y_base, items) in servicos:
        total = len(items)
        card_w = Inches(1.85)
        gap    = Inches(0.18)
        start_x = Inches(2.45)
        for j, (nome, badge, cor) in enumerate(items):
            x = start_x + j * (card_w + gap)
            add_rect(sl, x, y_base, card_w, Inches(1.0),
                     fill=C_BG_MID, line=cor, line_w=Pt(1.5))
            add_pill(sl, badge,
                     x + Inches(0.08), y_base + Inches(0.06),
                     Inches(0.52), Inches(0.3), cor, font_size=Pt(8))
            add_label(sl, nome,
                      x + Inches(0.08), y_base + Inches(0.38),
                      card_w - Inches(0.16), Inches(0.58),
                      font_size=Pt(9), color=C_WHITE)

# ---------------------------------------------------------------------------
# Slide 6 — Pipeline de Dados (fluxo)
# ---------------------------------------------------------------------------

def slide_pipeline(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Fluxo de Dados", Inches(0.7), Inches(0.25),
              Inches(9), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    etapas = [
        ("1", "Fontes\nBrutos",    "S3 Raw",        C_ACCENT2),
        ("2", "ETL\nGlue",         "Glue Job",      C_ACCENT),
        ("3", "Feature\nStore",    "SageMaker FS",  C_GREEN),
        ("4", "Treinar\nModelo",   "SageMaker",     C_GREEN),
        ("5", "Endpoint\nAPI",     "Inferência",    C_PURPLE),
        ("6", "Dashboard\nQuickSight", "Gestores",  C_ACCENT),
    ]

    n      = len(etapas)
    node_w = Inches(1.7)
    node_h = Inches(1.7)
    total_w = n * node_w + (n - 1) * Inches(0.6)
    start_x = (W - total_w) / 2
    y_node  = Inches(2.6)

    for i, (num, nome, sub, cor) in enumerate(etapas):
        x = start_x + i * (node_w + Inches(0.6))
        # círculo como quadrado arredondado
        shape = sl.shapes.add_shape(5, x, y_node, node_w, node_h)
        shape.line.color.rgb = cor
        shape.line.width = Pt(2)
        shape.fill.solid()
        shape.fill.fore_color.rgb = C_BG_MID
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = nome
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = C_WHITE

        add_label(sl, sub, x, y_node + node_h + Inches(0.12),
                  node_w, Inches(0.4),
                  font_size=Pt(9), color=cor, align=PP_ALIGN.CENTER)

        # badge número
        add_pill(sl, num,
                 x + node_w - Inches(0.42), y_node - Inches(0.18),
                 Inches(0.38), Inches(0.38), cor, font_size=Pt(10))

        # seta
        if i < n - 1:
            ax = x + node_w + Inches(0.08)
            ay = y_node + node_h / 2 - Inches(0.02)
            add_rect(sl, ax, ay, Inches(0.44), Inches(0.06), fill=cor)
            # pontinha da seta (triângulo simulado com pill)
            add_pill(sl, "▶",
                     ax + Inches(0.3), ay - Inches(0.14),
                     Inches(0.24), Inches(0.3), C_BG_DARK, fg=cor, font_size=Pt(9))

    # Retreinamento loop
    add_label(sl,
              "♻  EventBridge dispara retreinamento automático a cada semestre "
              "quando MAPE > 15%",
              Inches(1.5), Inches(5.8), Inches(10), Inches(0.55),
              font_size=Pt(12), color=C_GRAY, italic=True, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# Slide 7 — Modelo ML
# ---------------------------------------------------------------------------

def slide_modelo(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Modelo de Machine Learning", Inches(0.7), Inches(0.25),
              Inches(10), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    # Card principal — XGBoost
    add_rect(sl, Inches(0.65), Inches(1.1), Inches(5.8), Inches(5.8),
             fill=C_BG_MID, line=C_GREEN, line_w=Pt(2))
    add_label(sl, "⭐  XGBoost", Inches(0.9), Inches(1.25),
              Inches(5.2), Inches(0.6), font_size=Pt(20), bold=True, color=C_GREEN)
    add_label(sl, "Algoritmo principal", Inches(0.9), Inches(1.82),
              Inches(5.2), Inches(0.4), font_size=Pt(12), color=C_GRAY, italic=True)

    xg_detalhes = [
        ("Horizonte",         "2 semestres à frente"),
        ("Métrica principal", "MAPE (threshold 15%)"),
        ("Métricas adicionais","RMSE · MAE"),
        ("Instância treino",  "ml.m5.xlarge"),
        ("Retreinamento",     "A cada 180 dias (semestre)"),
    ]
    for i, (k, v) in enumerate(xg_detalhes):
        y = Inches(2.4) + i * Inches(0.78)
        add_label(sl, k, Inches(0.9),  y, Inches(2.3), Inches(0.55),
                  font_size=Pt(11), color=C_GRAY)
        add_label(sl, v, Inches(3.2),  y, Inches(3.1), Inches(0.55),
                  font_size=Pt(11), bold=True, color=C_WHITE)

    # Alternativas
    alts = [
        ("Prophet",  "Séries temporais com\nsazonalidade acadêmica",  C_ACCENT2),
        ("LSTM",     "Padrões sequenciais\ncomplexos e de longo prazo", C_PURPLE),
    ]
    add_label(sl, "Modelos Alternativos", Inches(7.1), Inches(1.1),
              Inches(5.5), Inches(0.5), font_size=Pt(16), bold=True, color=C_WHITE)
    for i, (nome, desc, cor) in enumerate(alts):
        y = Inches(1.75) + i * Inches(2.4)
        add_rect(sl, Inches(7.1), y, Inches(5.5), Inches(2.1),
                 fill=C_BG_MID, line=cor, line_w=Pt(2))
        add_label(sl, nome, Inches(7.3), y + Inches(0.15),
                  Inches(5.0), Inches(0.55), font_size=Pt(16), bold=True, color=cor)
        add_label(sl, desc, Inches(7.3), y + Inches(0.75),
                  Inches(5.0), Inches(1.0), font_size=Pt(12), color=C_LIGHT)

    # Trigger de retreinamento
    add_rect(sl, Inches(7.1), Inches(6.2), Inches(5.5), Inches(0.85),
             fill=C_BG_MID, line=C_ACCENT2, line_w=Pt(1.5))
    add_label(sl, "⚡  Trigger automático: MAPE > 15% → EventBridge → novo treino",
              Inches(7.25), Inches(6.35), Inches(5.2), Inches(0.55),
              font_size=Pt(10), color=C_ACCENT2)

# ---------------------------------------------------------------------------
# Slide 8 — Monitoramento e Operações
# ---------------------------------------------------------------------------

def slide_monitoramento(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Monitoramento & Operações", Inches(0.7), Inches(0.25),
              Inches(10), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    servicos_ops = [
        ("🔍", "SageMaker\nModel Monitor",
         "Detecção de Data Drift\nThreshold: 0.1 · Schedule: diário",
         C_GREEN),
        ("📊", "CloudWatch",
         "Alarmes: MAPE alto · Latência endpoint\nData Drift detectado",
         C_ACCENT),
        ("⏱", "EventBridge",
         "Agendador de retreinamento\nSchedule: a cada 180 dias",
         C_PURPLE),
        ("🚀", "Lambda Preprocessor",
         "Runtime: Python 3.12 · Timeout: 30s\nMemória: 512 MB",
         C_ACCENT2),
    ]

    for i, (icon, nome, desc, cor) in enumerate(servicos_ops):
        col = i % 2
        row = i // 2
        x = Inches(0.65) + col * Inches(6.2)
        y = Inches(1.2) + row * Inches(2.75)
        add_rect(sl, x, y, Inches(5.9), Inches(2.5),
                 fill=C_BG_MID, line=cor, line_w=Pt(2))
        add_label(sl, icon, x + Inches(0.15), y + Inches(0.15),
                  Inches(0.6), Inches(0.6), font_size=Pt(26))
        add_label(sl, nome, x + Inches(0.85), y + Inches(0.1),
                  Inches(4.8), Inches(0.7), font_size=Pt(14), bold=True, color=cor)
        add_label(sl, desc, x + Inches(0.2), y + Inches(1.05),
                  Inches(5.4), Inches(1.2), font_size=Pt(11), color=C_LIGHT)

# ---------------------------------------------------------------------------
# Slide 9 — Resultados Esperados / KPIs
# ---------------------------------------------------------------------------

def slide_kpis(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)
    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)

    add_label(sl, "Resultados Esperados", Inches(0.7), Inches(0.25),
              Inches(10), Inches(0.7), font_size=Pt(28), bold=True, color=C_WHITE)
    add_divider(sl, Inches(0.95))

    kpis = [
        ("< 15%",    "MAPE",             "Erro percentual\nabsoluto médio",          C_GREEN),
        ("2 sem.",   "Horizonte",        "Previsão antecipada\nde demanda",           C_ACCENT),
        ("180 dias", "Retreinamento",    "Ciclo automático\nde atualização",          C_ACCENT2),
        ("90%",      "Confiança",        "Intervalo de confiança\nnas previsões",     C_PURPLE),
    ]

    for i, (val, label, desc, cor) in enumerate(kpis):
        x = Inches(0.65) + i * Inches(3.15)
        y = Inches(1.2)
        # Card KPI
        add_rect(sl, x, y, Inches(2.9), Inches(2.5),
                 fill=C_BG_MID, line=cor, line_w=Pt(2))
        add_label(sl, val, x, y + Inches(0.3), Inches(2.9), Inches(1.0),
                  font_size=Pt(34), bold=True, color=cor, align=PP_ALIGN.CENTER)
        add_label(sl, label, x, y + Inches(1.25), Inches(2.9), Inches(0.45),
                  font_size=Pt(13), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
        add_label(sl, desc, x, y + Inches(1.75), Inches(2.9), Inches(0.65),
                  font_size=Pt(10), color=C_GRAY, align=PP_ALIGN.CENTER)

    # Padrões de nomenclatura
    add_label(sl, "Padrões de Nomenclatura AWS", Inches(0.65), Inches(3.95),
              Inches(11.8), Inches(0.5), font_size=Pt(16), bold=True, color=C_WHITE)
    add_divider(sl, Inches(4.4))

    padroes = [
        ("S3 Buckets",     "prevmatriculas-{env}-{tipo}",                   "ex: prevmatriculas-prod-raw"),
        ("Glue Jobs",      "job-prevmatriculas-{etapa}",                    "ex: job-prevmatriculas-feature-engineering"),
        ("Modelos SM",     "modelo-prevmatriculas-v{versao}",               "ex: modelo-prevmatriculas-v1"),
        ("Feature Group",  "prevmatriculas-features",                       "SageMaker Feature Store"),
    ]
    for i, (tipo, padrao, exemplo) in enumerate(padroes):
        y = Inches(4.55) + i * Inches(0.65)
        add_label(sl, tipo,    Inches(0.7),  y, Inches(2.0), Inches(0.55),
                  font_size=Pt(11), bold=True, color=C_GRAY)
        add_label(sl, padrao,  Inches(2.75), y, Inches(4.5), Inches(0.55),
                  font_size=Pt(11), color=C_ACCENT)
        add_label(sl, exemplo, Inches(7.35), y, Inches(5.5), Inches(0.55),
                  font_size=Pt(10), color=C_LIGHT, italic=True)

# ---------------------------------------------------------------------------
# Slide 10 — Encerramento
# ---------------------------------------------------------------------------

def slide_encerramento(prs):
    layout = prs.slide_layouts[6]
    sl = prs.slides.add_slide(layout)
    set_bg(sl, C_BG_DARK)

    add_rect(sl, 0, 0, W, Inches(0.08), fill=C_ACCENT)
    add_rect(sl, 0, Inches(7.42), W, Inches(0.08), fill=C_ACCENT2)
    gradient_rect(sl, 0, Inches(2.8), W, Inches(2.5), C_BG_MID, C_BG_DARK, angle=5400000)

    add_label(sl, "Obrigado!", Inches(0.8), Inches(1.2),
              Inches(11.5), Inches(1.2), font_size=Pt(58), bold=True,
              color=C_WHITE, align=PP_ALIGN.CENTER)

    add_label(sl,
              "Previsão de Demanda de Vagas por Curso  ·  Pipeline AWS End-to-End",
              Inches(0.8), Inches(2.6), Inches(11.5), Inches(0.6),
              font_size=Pt(18), color=C_ACCENT, align=PP_ALIGN.CENTER)

    stack = [
        "☁  S3  ·  Glue  ·  SageMaker  ·  Lambda",
        "📊  CloudWatch  ·  EventBridge  ·  QuickSight",
        "🤖  XGBoost  ·  MAPE < 15%  ·  2 semestres de horizonte",
    ]
    for i, linha in enumerate(stack):
        add_label(sl, linha,
                  Inches(0.8), Inches(3.7) + i * Inches(0.62),
                  Inches(11.5), Inches(0.55),
                  font_size=Pt(14), color=C_LIGHT, align=PP_ALIGN.CENTER)

    add_label(sl, "Workshop Arquiteturas AWS  ·  Kiro IDE  ·  2026",
              Inches(0.8), Inches(6.85), Inches(11.5), Inches(0.4),
              font_size=Pt(11), color=C_GRAY, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    print("Gerando slides...")
    slide_capa(prs);            print("  ✓ Slide 1 — Capa")
    slide_problema(prs);        print("  ✓ Slide 2 — Problema e Solução")
    slide_dados(prs);           print("  ✓ Slide 3 — Fontes de Dados")
    slide_features(prs);        print("  ✓ Slide 4 — Feature Engineering")
    slide_arquitetura(prs);     print("  ✓ Slide 5 — Arquitetura AWS")
    slide_pipeline(prs);        print("  ✓ Slide 6 — Fluxo de Dados")
    slide_modelo(prs);          print("  ✓ Slide 7 — Modelo ML")
    slide_monitoramento(prs);   print("  ✓ Slide 8 — Monitoramento")
    slide_kpis(prs);            print("  ✓ Slide 9 — KPIs e Padrões")
    slide_encerramento(prs);    print("  ✓ Slide 10 — Encerramento")

    out = os.path.join(os.path.dirname(__file__), "previsao-matriculas.pptx")
    prs.save(out)
    print(f"\n✅  Salvo em: {out}")

if __name__ == "__main__":
    main()
